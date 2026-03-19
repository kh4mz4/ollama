from pathlib import Path
from typing import List

import openpyxl
from llama_index.core import (
    Document,
    SimpleDirectoryReader,
)

from system.config import SUPPORTED_EXTS, DOCS_DIR


def docs_exist(folder: Path) -> bool:
    if not folder.exists() or not folder.is_dir():
        return False
    for p in folder.rglob("*"):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS and p.stat().st_size > 0:
            return True
    return False


def xlsx_to_documents(
        xlsx_path: Path,
        rows_per_doc: int = 120,
) -> List[Document]:
    """
    Converts XLSX into multiple Documents by chunking rows.
    Each chunk includes a header row (first non-empty row) if present.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    docs: List[Document] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Read rows into list[str] (TSV lines)
        all_lines: List[str] = []
        for row in ws.iter_rows(values_only=True):
            if row is None:
                continue
            cells = ["" if c is None else str(c) for c in row]
            if all(c.strip() == "" for c in cells):
                continue
            all_lines.append("\t".join(cells))

        if not all_lines:
            continue

        header = all_lines[0]
        data_lines = all_lines[1:] if len(all_lines) > 1 else []

        if not data_lines:
            # sheet with only 1 row
            text = (
                f"File: {xlsx_path.name}\n"
                f"Sheet: {sheet_name}\n"
                f"---\n"
                f"{header}\n"
            )
            docs.append(
                Document(
                    text=text,
                    metadata={
                        "source": str(xlsx_path),
                        "file_name": xlsx_path.name,
                        "sheet": sheet_name,
                        "row_range": "1-1",
                        "type": "xlsx",
                    },
                )
            )
            continue

        # Chunk rows
        start_row = 2  # because header is row 1 in our text representation
        for i in range(0, len(data_lines), rows_per_doc):
            chunk = data_lines[i: i + rows_per_doc]
            chunk_start = start_row + i
            chunk_end = chunk_start + len(chunk) - 1

            text = (
                    f"File: {xlsx_path.name}\n"
                    f"Sheet: {sheet_name}\n"
                    f"Rows: {chunk_start}-{chunk_end}\n"
                    f"---\n"
                    f"{header}\n"
                    + "\n".join(chunk)
            )

            docs.append(
                Document(
                    text=text,
                    metadata={
                        "source": str(xlsx_path),
                        "file_name": xlsx_path.name,
                        "sheet": sheet_name,
                        "row_range": f"{chunk_start}-{chunk_end}",
                        "type": "xlsx",
                    },
                )
            )

    return docs


def load_all_documents() -> List[Document]:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)

    if not docs_exist(DOCS_DIR):
        print(f"❌ No documents found in: {DOCS_DIR}")
        print("➡️  Put files there (pdf/txt/docx/md/py/json/csv/xlsx) then run again.")
        raise SystemExit(1)

    non_xlsx: List[Path] = []
    xlsx_files: List[Path] = []

    for p in DOCS_DIR.rglob("*"):
        if not p.is_file():
            continue
        if p.suffix.lower() not in SUPPORTED_EXTS:
            continue
        if p.stat().st_size == 0:
            continue
        if p.suffix.lower() == ".xlsx":
            xlsx_files.append(p)
        else:
            non_xlsx.append(p)

    docs: List[Document] = []

    if non_xlsx:
        # Let LlamaIndex read PDFs, docx, text, code, etc.
        docs.extend(SimpleDirectoryReader(input_files=[str(p) for p in non_xlsx]).load_data())

    for xf in xlsx_files:
        docs.extend(xlsx_to_documents(xf, rows_per_doc=120))

    return docs
